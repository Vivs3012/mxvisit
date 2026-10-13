#!/usr/bin/env python3
"""Peru Consumer & Rewards deck — Top Tiendas - TaDa.pptx template."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/Users/zeuser/Desktop/Top Tiendas - TaDa.pptx")
OUT = ROOT / "export" / "PE-Consumer-Rewards-Executive-Peru.pptx"
PHOTOS = ROOT / "export" / "workshop-photos"

PURPLE = RGBColor(0x9c, 0x41, 0xff)
WHITE = RGBColor(0xff, 0xff, 0xff)
GRAY = RGBColor(0xcc, 0xcc, 0xcc)
TITLE_FONT = "Avantt TRIAL ExtraBold"
BODY_FONT = "Avantt TRIAL"
TITLE_BOX = (0.705, 0.67, 11.5, 1.35)


def delete_all_slides(prs):
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]


def style_run(run, size_pt, bold=False, color=WHITE, font=BODY_FONT):
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color


def set_title_text(shape, lines):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    for i, line in enumerate(lines):
        if i > 0:
            p = tf.add_paragraph()
        r = p.add_run()
        r.text = line.upper()
        style_run(r, 44, True, PURPLE, TITLE_FONT)


def add_title_shape(slide, lines):
    left, top, w, h = TITLE_BOX
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    set_title_text(box, lines)


def add_subtitle(slide, text, top=1.52):
    if not text:
        return
    box = slide.shapes.add_textbox(Inches(0.705), Inches(top), Inches(11.8), Inches(0.5))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    style_run(r, 13, False, GRAY)


def add_body_box(slide, paragraphs, left=0.705, top=2.05, width=11.85, height=5.15, body_pt=13.5, label_pt=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, block in enumerate(paragraphs):
        if isinstance(block, tuple):
            label, bullets = block
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(8)
            r = p.add_run()
            r.text = label
            style_run(r, label_pt, True, PURPLE, TITLE_FONT)
            for b in bullets:
                p = tf.add_paragraph()
                p.space_after = Pt(3)
                r = p.add_run()
                r.text = f"• {b}"
                style_run(r, body_pt, False, WHITE)
        else:
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(5)
            r = p.add_run()
            r.text = block
            style_run(r, body_pt, False, WHITE)


def new_black_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[0])


def add_workshop_slide(prs):
    slide = new_black_slide(prs)
    add_title_shape(slide, ["Growth", "Workshop"])
    add_subtitle(slide, "Sep 2, 2026 · BeerHub, Lima · Consumer + Rewards (shared board, separate votes)")
    add_body_box(
        slide,
        [
            ("How it ran", [
                "Same day as OPS workshop — 3h One Year Plan 2027 with Product visit (Brazil/Mexico).",
                "Consumer and Rewards ideated on a shared board; each track presented priorities and voted top 4.",
                "Part 2: future vision (personalization, omnichannel loyalty, football occasion, social commerce).",
                "Closing: El Momento Inesperado — hospitality surprises (football kits, birthdays, Club first order).",
            ]),
            ("Peru context", [
                "National awareness ~9% (Kantar) — growth and loyalty must earn trust, not only promos.",
                "Club TaDa launched during visit week; RGB ~44% of orders but low education on returnables.",
                "Consumer reliability gaps (map, address) overlap Service Level — same fixes unlock both.",
            ]),
        ],
        left=0.705,
        top=2.05,
        width=5.55,
        height=5.0,
        body_pt=13,
    )
    photos = sorted(PHOTOS.glob("*.png"))
    if len(photos) >= 2:
        slide.shapes.add_picture(str(photos[0]), Inches(6.35), Inches(0.95), width=Inches(3.15))
        slide.shapes.add_picture(str(photos[1]), Inches(9.55), Inches(0.95), width=Inches(3.15))


def add_meetings_slide(prs):
    slide = new_black_slide(prs)
    add_title_shape(slide, ["Field", "Meetings"])
    add_subtitle(
        slide,
        "Aug 31 → Sep 2 → Sep 3 · Context, stores, workshop votes, then Consumer product alignment",
    )
    add_body_box(
        slide,
        [
            ("Aug 31 · Context meeting (Lima office)", [
                "Peru Growth/Commercial (Ulises, Eloy, Alex, acquisition & retention leads) + visiting product — same morning as POC visits.",
                "Growth engine: 102k new customers target; healthy acquisition (no new-user coupon); TikTok ~40% of paid, Meta ~30%; homemade creative in Lima.",
                "Retention gap: RGB M1→M2 31.3% vs Marketplace-only 20.7%; only 27 of 100 new users buy returnables — education + price, not awareness alone.",
                "Club TaDa launch that week — points ladder (~2% effective), 10% Loyalty Rate year target; short term coupons, mid term Backus experiences (football, concerts).",
                "Awareness Kantar: TaDa 9% vs Rappi 41% / PedidosYa 37%; moats = returnables + cold drinks vs aggregators.",
                "Portfolio: monetary coupons steered 100% to beer so Marketplace + Take Rate does not go negative for POCs; brand pages ready locally but create-in-admin still engineering-only.",
                "Product levers on table: copy Brazil location UX (config); checkout address-review modal just launched; ZD Libri + Consumer Room for quant/qual research.",
            ]),
            ("Aug 31 · Store visits — La Marina, Merino, San Pedro", [
                "3/3 POCs: consumer map pin stuck (WhatsApp is real-time layer) — same theme as workshop #1 and Sep 3 P1.",
                "Search misses brand/format; deposit-bottle (“tengo/no tengo”) disputes; clients cancel repeatedly to route to preferred POC (store name visible).",
                "Consumer promo in app without BEES alert — stores cannot pre-stock for campaigns; hurts fill rate and campaign ROI.",
            ]),
            ("Sep 2 · Growth workshop (BeerHub) — summary", [
                "3h One Year Plan 2027; Consumer + Rewards shared board, separate top-4 votes (detail + photos on next slide).",
                "Consumer: map, MMP/data, payment gateways, wallet UX. Rewards: Braze enrichment, flash coupons, BEES admin, anti-fraud.",
                "Future vision + El Momento Inesperado hospitality (kits, birthdays, first Club order).",
            ]),
            ("Sep 3 · Consumer vision with Growth", [
                "Victoria’s squad walked 12+ 2026 features (multishowcases, Buy Again 61% conversion, address flow, Live Activities Sep 1, brand pages, Turbo pilot).",
                "Eloy’s team listed 18 prioritized pains — 6 P1 consumer-centric (RGB education, map, partial order/stock, CX bot, search, wrong address / no-answer).",
                "Explicit bridges to Rewards: coupon wallet lifecycle (P2), Club TaDa ordering in admin, Braze cache (kill app), Fish Promo personalization (+6% GMV MX reference).",
                "Next steps captured: feature prioritization for Peru (not all at once), Braze event audit with Bruno, lifecycle deck for Club migration, returnables brand page for next campaign.",
            ]),
        ],
        body_pt=11.5,
        label_pt=14,
    )


def build():
    if not TEMPLATE.is_file():
        raise SystemExit(f"Template not found: {TEMPLATE}")

    prs = Presentation(str(TEMPLATE))
    delete_all_slides(prs)

    slide = new_black_slide(prs)
    add_title_shape(slide, ["Peru", "Consumer & Rewards"])
    add_body_box(
        slide,
        [
            "Executive learnings for Product — Growth & loyalty tracks only.",
            "",
            "Sources: Aug 31 context + store visits · Sep 2 workshop · Sep 3 Consumer vision · Club TaDa launch week.",
            "",
            "Scope: not Service Level (Seller/Courier/Admin/CX) — see companion deck.",
            "Headline: reliability and data foundations before scale — map, MMP/Braze, wallet UX, then tiers/omnichannel.",
        ],
        top=2.05,
        body_pt=15,
    )

    add_meetings_slide(prs)

    add_workshop_slide(prs)

    slide = new_black_slide(prs)
    add_title_shape(slide, ["Consumer"])
    add_subtitle(slide, "Workshop Sep 2 + Vision Sep 3 · 18 Growth pains prioritized with commercial team")
    add_body_box(
        slide,
        [
            ("What we learned", [
                "Live courier map is #1 ask — pin rarely moves; ETA reduced to ranges; overlaps seller map failure (3/3 stores).",
                "Data stack not homologated: AppsFlyer/Segment gaps block attribution, Braze triggers, and paid media optimization.",
                "Payments ~3% online — single gateway (Niubiz) blocks BIN/bank campaigns like Mercado Pago.",
                "Returnables/RGB (~44% orders): consumers don't understand exchange rules; inter-brand swap unclear.",
                "P1 gaps: search/discovery limits, CX bot not tropicalized (~45% deflection but stale AI), partial order blocked without seller stock quantity, consumer no-answer (#2 cancel driver), wrong address ~18%.",
                "Wallet/coupons: manual lifecycle clusters; Club TaDa ordering helps but flash UX and Braze cache (kill app) hurt campaigns.",
            ]),
            ("Workshop top 4 votes", [
                "#1 Real-time courier tracking on map.",
                "#2 SDK + MMP integration & data homologation (close AppsFlyer order gap).",
                "#3 Payment gateways — BIN / bank partnerships (BBVA, BCP).",
                "#4 Coupon wallet UX & ordering (Club TaDa without burying other coupons).",
            ]),
            ("Shipping / near term (Sep 3)", [
                "Multishowcases (Sep 15 target) · Buy Again 61% order conversion · new address flow (CO live).",
                "Live Activities on lock screen (Sep 1) · search typo fallback · brand pages for returnables education.",
            ]),
        ],
        body_pt=12.5,
        label_pt=15,
    )

    slide = new_black_slide(prs)
    add_title_shape(slide, ["Rewards"])
    add_subtitle(slide, "Workshop Sep 2 + Club TaDa launch week · dependency-heavy stack before tiers")
    add_body_box(
        slide,
        [
            ("What we learned", [
                "Club TaDa launched during visit — ~85% active users enrolled, mostly B tier; tiers not live yet.",
                "Heavy reliance on monetary coupons; want Braze enrichment — points for surveys, channels, data (vote #1).",
                "Flash burn coupons with countdown not built; wallet visuals and coupon images lag competitors.",
                "Rewards admin outside BEES — CAO list downloads, manual anti-fraud for bodegueros, raffle winners manual.",
                "SessionM ↔ Bees One catalog sync friction; batch points and custom events need hardening.",
                "Future vision voted: football occasion earn/burn + API, omnichannel redeem (Tambo/OXXO test), tiers 3–4 yr horizon.",
                "National awareness 9% — loyalty must pair with reliability (Consumer map/address) and CRM data quality.",
            ]),
            ("Workshop top 4 votes", [
                "#1 Club TaDa × Braze — points for enrichment, surveys, channel opt-in.",
                "#2 Flash coupon (burn) with time-boxed countdown in wallet.",
                "#3 Rewards admin inside BEES — coupons, missions, rules in one place.",
                "#4 Automated anti-fraud — bodeguero patterns, abnormal earn/burn.",
            ]),
            ("Now vs later", [
                "Now: wallet UX + flash mechanics with Club admin · Braze triggers · BEES admin centralization.",
                "Gate: homologate events before scaling personalized earn/burn.",
                "Later: tiers, omnichannel redemption, football API, premium hospitality experiences.",
            ]),
        ],
        body_pt=12.5,
        label_pt=15,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()

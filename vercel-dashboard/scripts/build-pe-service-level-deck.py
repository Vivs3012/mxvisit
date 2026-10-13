#!/usr/bin/env python3
"""Peru Service Level deck — styled from Top Tiendas - TaDa.pptx template."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/Users/zeuser/Desktop/Top Tiendas - TaDa.pptx")
OUT = ROOT / "export" / "PE-Service-Level-Executive-Peru.pptx"
PHOTOS = ROOT / "export" / "workshop-photos"

PURPLE = RGBColor(0x9c, 0x41, 0xff)
WHITE = RGBColor(0xff, 0xff, 0xff)
GRAY = RGBColor(0xcc, 0xcc, 0xcc)

TITLE_FONT = "Avantt TRIAL ExtraBold"
BODY_FONT = "Avantt TRIAL"

# Match template slide 3 text box
BODY_BOX = (0.705, 2.53, 5.96, 4.85)
TITLE_BOX = (0.705, 0.67, 11.5, 1.35)


def delete_all_slides(prs):
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]


def style_run(run, size_pt, bold=False, color=WHITE, font=BODY_FONT):
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color


def set_title_text(shape, lines):
    """lines: list of strings; rendered like template with line breaks."""
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    for i, line in enumerate(lines):
        if i > 0:
            p = tf.add_paragraph()
        r = p.add_run()
        r.text = line.upper()
        style_run(r, 44, True, PURPLE, TITLE_FONT)
        p.alignment = PP_ALIGN.LEFT


def add_title_shape(slide, lines):
    left, top, w, h = TITLE_BOX
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    set_title_text(box, lines)
    return box


def add_subtitle(slide, text, top=1.52):
    if not text:
        return
    box = slide.shapes.add_textbox(Inches(0.705), Inches(top), Inches(11.8), Inches(0.5))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    style_run(r, 13, False, GRAY)


def add_body_box(
    slide,
    paragraphs,
    left=None,
    top=None,
    width=None,
    height=None,
    body_pt=15,
    label_pt=17,
):
    l, t, w, h = BODY_BOX
    if left is not None:
        l, t, w, h = left, top, width, height
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, block in enumerate(paragraphs):
        if isinstance(block, tuple):
            label, bullets = block
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(8)
            r = p.add_run()
            r.text = label
            style_run(r, label_pt, True, PURPLE, TITLE_FONT)
            for b in bullets:
                p = tf.add_paragraph()
                p.level = 0
                p.space_after = Pt(3)
                r = p.add_run()
                r.text = f"• {b}"
                style_run(r, body_pt, False, WHITE)
        else:
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(5)
            r = p.add_run()
            r.text = block
            style_run(r, body_pt, False, WHITE)


def new_black_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[0])  # Slide negro


def build():
    if not TEMPLATE.is_file():
        raise SystemExit(f"Template not found: {TEMPLATE}")

    prs = Presentation(str(TEMPLATE))
    delete_all_slides(prs)

    # Cover
    slide = new_black_slide(prs)
    add_title_shape(slide, ["Peru", "Service Level"])
    add_body_box(
        slide,
        [
            "Executive learnings for Product — Seller · Courier · Admin · CX",
            "",
            "Sources: Aug 31 store visits (3 Lima POCs) · Sep 1 Service Level with OPS + Product · Sep 2 OPS workshop (BeerHub).",
            "",
            "Scope: operations & service level only — not Consumer or Rewards tracks.",
            "Peru context: +36% orders Jan–Jul YoY, ~96% SLA, ~18% cancellations from wrong address.",
            "Shared product bets: Ecosystem Communication (weeks) · real-time inventory (strategic) · Braze for supervisor→store comms.",
        ],
        left=0.705,
        top=2.05,
        width=11.8,
        height=4.8,
        body_pt=15,
    )

    # Workshop
    slide = new_black_slide(prs)
    add_title_shape(slide, ["Peru OPS", "Workshop"])
    add_body_box(
        slide,
        [
            ("How it ran", [
                "Kick-off → challenge board (real ops examples) → group prioritization vote → future vision for 2027.",
                "Dedicated Operations board — stores, couriers, CX, admin (27.CX.ADMI).",
                "Part 1: 25+ operational pains on stickies. Part 2: co-created bets (scheduled orders, copilot, BEES→TADA, vending).",
            ]),
            ("Who participated", [
                "Peru OPS: supervisors, stores, couriers, CIEX/CX, commercial admin.",
                "Visiting Product & Design (Brazil/Mexico); Consumer and Rewards on a parallel board the same day.",
            ]),
            ("Top 4 votes — Operations", [
                "#1 Real-time KPI visibility — stock, SLA, availability, hot CRM SKUs; alerts before SLA breaks.",
                "#2 Automation + AI copilot for POC admin — coverage, polygons, grouping with SLA guardrails.",
                "#3 Claims resolution — CX, store & consumer on one path; fast credit like market leaders, not parallel email.",
                "#4 SKU restrictions on coupons — fewer promo errors, cancellations, and unfair seller KPI penalties.",
            ]),
        ],
        left=0.705,
        top=2.05,
        width=5.55,
        height=5.0,
        body_pt=13,
        label_pt=16,
    )
    photos = sorted(PHOTOS.glob("*.png"))
    if len(photos) >= 2:
        slide.shapes.add_picture(str(photos[0]), Inches(6.35), Inches(0.95), width=Inches(3.15))
        slide.shapes.add_picture(str(photos[1]), Inches(9.55), Inches(0.95), width=Inches(3.15))
    cap = slide.shapes.add_textbox(Inches(6.35), Inches(6.55), Inches(6.3), Inches(0.4))
    cr = cap.text_frame.paragraphs[0].add_run()
    cr.text = "Sep 2, 2026 · BeerHub, Lima · One Year Plan 2027"
    style_run(cr, 12, False, GRAY)

    actors = [
        (
            ["Seller"],
            "TaDa Seller app + Web Tiendas (WebPO) · validated in 3/3 Lima POCs",
            [
                ("What we learned", [
                    "Map does not update: courier pin stays at the store; on peak only 1–2 of 15 riders visible — stores track via WhatsApp (La Marina ~5–10% visibility).",
                    "Inventory is on/off only (no quantity). Combos stay active when one SKU is off; search misses brands/formats; orders still arrive on disabled SKUs.",
                    "Order alert sound fails or plays on the wrong device (shared login: office PC vs shop floor) → missed and expired orders.",
                    "No in-app chat with the customer; deposit-bottle disputes hurt ratings; cannot assign a 2nd order to a courier already en route (La Marina #2 ask).",
                    "WebPO is valued (demand by hour, courier planning, weekly payments) but consumer promos often do not alert the store to pre-stock.",
                ]),
                ("Priorities for Product", [
                    "Stock quantity + partial order when one SKU is out (today: full cancel).",
                    "Ecosystem Communication (seller ↔ consumer ↔ courier) — shipping in weeks.",
                    "Actionable KPIs in-app; Top Tiendas off Monday Excel/WhatsApp.",
                    "Reliable push + sound + session on shared store accounts.",
                ]),
            ],
        ),
        (
            ["Courier"],
            "Fleet & last mile · driven by address quality, wait time, and proof of delivery",
            [
                ("What we learned", [
                    "Wrong address drives ~18% of cancellations (saved addresses + checkout UX); INAP review launched before the visit.",
                    "Customer no-answer at the door: no standard wait time → courier/seller cancel; major ticket and rating driver.",
                    "Sellers dispatch blind when the map freezes; WhatsApp is the real-time layer (no audit trail).",
                    "Returnable / heavy promos create unsafe stacking on bikes; store phone sometimes exposed to couriers.",
                    "“Delivered but not received” is low volume but painful; delivery PIN exists globally — Peru store-level pilot.",
                ]),
                ("Priorities for Product", [
                    "Address UX (checkout review + location patterns from CO/BR tests).",
                    "Timed courier ↔ consumer messages (5-min window; MX pilot −30% courier cancels).",
                    "Live map shared with seller (and CX) — same layer, one fix.",
                    "Roll out delivery confirmation code per POC.",
                ]),
            ],
        ),
        (
            ["Admin /", "OPS"],
            "Commercial back-office + network supervision — hidden load on service level",
            [
                ("What we learned", [
                    "No real-time network cockpit: at rush, supervisors cannot see offline stores, expired orders, or new CRM SKUs not activated — Power BI T+1 + WhatsApp.",
                    "Promo/discount ops: 4 discount types × DDC sheets → 8+ templates updated twice weekly; flash promos require delete/re-upload; shelf order fixed manually each time.",
                    "Polygons edited street-by-street; no neighbor POC view; dangerous zones not excludable (BEES has this, TaDa does not).",
                    "Braze (notification center) planned to replace supervisor WhatsApp blasts; Admin squad capacity blocks the ops dashboard.",
                    "~20% of new POCs close within 3–4 months — gold/silver/bronze ranking still lives outside the product.",
                ]),
                ("Priorities for Product", [
                    "KPI cockpit per POC (resume OPS Power BI inside the platform; weekend on-call).",
                    "Braze: targeted store comms with read metrics.",
                    "Promo admin: change one campaign without full spreadsheet re-upload.",
                    "WebPO: keep weekly VLC/credit notes; close gap on per-order export or stores stay on legacy tools.",
                ]),
            ],
        ),
        (
            ["CX"],
            "CIEX · Zendesk · triangle between store, consumer, and support",
            [
                ("What we learned", [
                    "Claims path is slow: Zendesk → email to store → days for evidence; CX, store, and consumer are not on one thread.",
                    "Support cannot cancel after the courier has the order — parallel email/Teams.",
                    "Small-claim coupons deflect ~40% of tickets (working); large cases (full case damage) still need the store.",
                    "CX has no live courier map — same blind spot as consumers; drives “where is my order?” volume.",
                    "Tax docs: customer puts RUC in delivery notes; store does not see it → ticket days later; checkout checkbox missing.",
                    "OPS workshop vote #3: resolve claims like market leaders (fast credit), not parallel email.",
                ]),
                ("Priorities for Product", [
                    "Faster claims triangle + clearer roles (when store vs automated coupon).",
                    "cx-courier-map-visibility — align with seller map fix.",
                    "Ecosystem Communication reduces issues that become tickets.",
                    "Checkout invoice type + RUC; delivery PIN for dispute closure.",
                ]),
            ],
        ),
    ]

    for title_lines, subtitle, sections in actors:
        slide = new_black_slide(prs)
        add_title_shape(slide, title_lines)
        add_subtitle(slide, subtitle)
        add_body_box(
            slide,
            sections,
            left=0.705,
            top=2.05,
            width=11.85,
            height=5.15,
            body_pt=13.5,
            label_pt=16,
        )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
